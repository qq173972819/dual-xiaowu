from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"E:\electron\screen\inksight-main\inksight-main\xiaowu-s3-open-platform")
OUT = ROOT / "output" / "whitepaper"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "小悟AI墨水屏产品白皮书.docx"
ARCH = OUT / "技术架构.png"

NAVY = "14324A"
TEAL = "2A8C82"
PALE = "EAF3F1"
GRAY = "5D6670"
LIGHT = "F4F6F7"

def font(size, bold=False, color="000000", name="Microsoft YaHei"):
    def apply(run):
        run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
        run.font.name = name
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    return apply

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def set_cell_border(cell, color="D9D9D9", size="6"):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), size); el.set(qn("w:color"), color)

def margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=OxmlElement("w:"+m); node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa"); tcMar.append(node)

def make_architecture():
    W,H=1600,1050
    im=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(im)
    try:
        fb=ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc",38)
        fm=ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc",27)
        fs=ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc",23)
    except:
        fb=fm=fs=None
    d.text((70,45),"小悟 AI 墨水屏技术架构",font=fb,fill="#14324A")
    layers=[
        ("应用与场景层",["家庭看板","学习陪练","健康管理","信息简报","文化内容","商业展示"],"#EAF3F1"),
        ("AI 与内容服务层",["大模型接入","语音识别与合成","内容生成","模板编排","图像处理","规则与日程"],"#EDF3F8"),
        ("云端开放平台层",["设备管理","模式广场","内容后台","刷新计划","用户与权限","数据分析"],"#F4F6F7"),
        ("连接与协议层",["Wi-Fi","Bluetooth","NFC","HTTPS / MQTT","Home Assistant","OTA / Web Flash"],"#FFF5E9"),
        ("设备与硬件层",["ESP32-C3","ESP32-S3","麦克风 / 扬声器","电池与休眠","1.54–10 英寸","双色至 E6 全彩"],"#F7EFEA"),
    ]
    y=130
    for title,items,fill in layers:
        d.rounded_rectangle((70,y,1530,y+150),radius=22,fill=fill,outline="#CBD3D8",width=3)
        d.text((95,y+22),title,font=fm,fill="#14324A")
        x=375
        for item in items:
            d.rounded_rectangle((x,y+28,x+175,y+120),radius=14,fill="white",outline="#B9C5CB",width=2)
            bb=d.textbbox((0,0),item,font=fs); tw=bb[2]-bb[0]; th=bb[3]-bb[1]
            d.text((x+(175-tw)/2,y+74-th/2),item,font=fs,fill="#33434D")
            x+=186
        y+=172
    d.text((70,1000),"架构原则：端云协同 · 低功耗常显 · 多尺寸适配 · 模式可扩展 · 软硬件一体化交付",font=fs,fill="#5D6670")
    im.save(ARCH,quality=95)

doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
styles["Normal"].font.name="Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); styles["Normal"].font.size=Pt(10); styles["Normal"].font.color.rgb=RGBColor.from_string("26343D")
styles["Title"].font.name="Microsoft YaHei"; styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); styles["Title"].font.size=Pt(28); styles["Title"].font.bold=True; styles["Title"].font.color.rgb=RGBColor(0,0,0)
for n,size in (("Heading 1",22),("Heading 2",14)):
    s=styles[n]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)

def add_text(text,size=10,bold=False,color="26343D",align=None,space_after=7):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(space_after); p.paragraph_format.line_spacing=1.35
    if align is not None: p.alignment=align
    r=p.add_run(text); font(size,bold,color)(r); return p

def title_page(section, title, subtitle=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); r=p.add_run(section); font(9,True,TEAL)(r)
    p=doc.add_paragraph(style="Heading 1"); p.paragraph_format.space_after=Pt(12); p.add_run(title)
    if subtitle: add_text(subtitle,11,False,GRAY,space_after=14)

def page(): doc.add_page_break()

def picture(path,width,caption=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5); p.add_run().add_picture(str(path),width=Inches(width))
    if caption: add_text(caption,8,False,GRAY,WD_ALIGN_PARAGRAPH.CENTER,8)

def bullets(items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(4); p.paragraph_format.left_indent=Inches(.18); r=p.add_run(item); font(10)(r)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_cell_shading(c,NAVY); set_cell_border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); font(9,True,"FFFFFF")(r)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; set_cell_border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: set_cell_shading(c,"F5F8FA")
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run(val); font(8.6)(r)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

logo=ROOT/"logo1-xiaowu-ai.png"
hero=ROOT/"firmware/s3/docs/xiaowu/voice-note-product.jpg"
e6=ROOT/"firmware/s3/docs/xiaowu/seven-inch-e6-display.jpg"
board=ROOT/"firmware/s3/docs/xiaowu/dual-screen-mainboard.jpg"
nfc=Path(r"C:\Users\upisq\AppData\Local\Temp\browser-use\assets\856ec2cf-e77d-4e00-9a4b-eb9e94fa0919\93d86fee8eac0ee1.jpg")
c3=ROOT/"firmware/s3/docs/v1/xmini-c3.jpg"
daily=ROOT/"services/legacy-api/docs/design/daily-800x480-after.png"

# Cover
picture(logo,1.25)
p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(8); p.add_run("小悟 AI 墨水屏产品白皮书")
add_text("低功耗常显设备与 AI 内容生态的一体化产品体系",13,False,GRAY,WD_ALIGN_PARAGRAPH.CENTER,16)
picture(hero,6.5)
add_text("沄墨电子科技(四川)有限责任公司",11,True,NAVY,WD_ALIGN_PARAGRAPH.CENTER,3)
add_text("2026 年 10 月  产品与技术概览",9,False,GRAY,WD_ALIGN_PARAGRAPH.CENTER,0)

# 2
page(); title_page("01  白皮书摘要","让信息安静地留在眼前","小悟以电子纸为长期可见的信息载体，把设备、内容、AI、后台和家庭物联网连接为同一套可扩展系统。")
table(["核心价值","产品实现","客户收益"],[["低功耗常显","电子纸仅在刷新时主要耗电，断电后仍保留画面","减少发光屏依赖，适合桌面与空间常驻"],["AI 内容生成","接入 DeepSeek、通义千问、Kimi 等模型能力","让天气、时间和场景转化为个性内容"],["软硬件一体化","设备、固件、后台、模式与开放平台协同","从样机验证延伸到批量交付和定制"],["多形态覆盖","1.54 至 10 英寸，静默、语音、挂件与全彩形态","覆盖个人、家庭、教育及商业展示"]],[1.45,2.35,2.65])
add_text("产品体系结论",12,True,NAVY,space_after=5)
add_text("小悟并非单一电子纸硬件，而是一套围绕常显信息、AI 内容和设备管理构建的产品生态。统一后台可管理设备分组、内容模板与刷新计划，同一内容可面向不同尺寸和色彩能力进行适配。")
bullets(["4 类核心产品系列，5 类可选显示尺寸，覆盖双色、三色、四色、六色及 E6 全彩。","C3 系列侧重低功耗静默展示，S3 系列扩展语音交互与学习陪伴。","支持在线刷机、设备配置、无设备预览、模式广场和自定义内容模式。"])

# 3
page(); title_page("02  产品目录","四类产品覆盖从桌面到空间展示","量产产品、研发版本与定制方案，以最终交付规格为准。")
table(["产品系列","典型规格","交互方式","主要场景"],[["ESP32-C3 静默摆件","4.2 / 5.83 英寸，三色 / 四色","Wi-Fi，按计划刷新","桌面信息、家庭看板、艺术内容"],["ESP32-S3 语音陪伴","单屏 / 双屏，麦克风与扬声器","语音、按键、菜单","教育陪练、阅读、健康管理"],["NFC / 蓝牙挂件","1.54 英寸，四色","NFC + Bluetooth","随身挂件、冰箱贴、身份与提示"],["大尺寸全彩墨水屏","7 / 10 英寸，E6 全彩，可扩展语音","Wi-Fi / 语音 / 后台","商业展示、艺术画框、空间信息"]],[1.35,2.1,1.55,1.8])
picture(e6,4.8,"7 英寸 E6 全彩电子纸实机示例")
add_text("产品选型建议",12,True,NAVY,space_after=4)
add_text("需要最低功耗与安静常显时优先选择 C3；需要语音与复杂交互时选择 S3；需要轻便近场更新时选择 NFC / 蓝牙挂件；需要高视觉表现和空间展示时选择大尺寸全彩系列。")

# 4
page(); title_page("03  ESP32 C3 静默摆件","桌面上的多场景 AI 搭档","在线刷机、设备配置、无设备预览与模式广场共同构成低门槛使用路径。")
picture(c3,5.1,"ESP32-C3 开发形态示例")
table(["能力","说明"],[["AI 驱动","根据天气、时间与场景生成个性化内容"],["电子墨水","类纸显示、不主动发光，适合长时间常驻"],["模式生态","内置 24 类内容模式，支持自定义与分享"],["低功耗设计","RISC-V 架构、Wi-Fi、深度休眠与电池供电"],["开发支持","Arduino、ESP-IDF、PlatformIO 与在线刷机"]],[1.55,4.9])
add_text("官网公开信息显示，C3 开发板方案可采用 LiFePO4 锂电池与 Deep Sleep 模式，并以开源硬件和可制作性为设计方向。")

# 5
page(); title_page("04  ESP32 S3 语音陪伴终端","会开口的墨水屏陪伴每一次练习","单屏与双屏两种形态，将低干扰显示与 AI 语音交互结合。")
picture(hero,5.9,"语音便签与语音交互产品形态示意")
table(["模块","典型能力"],[["语音交互","麦克风、扬声器、AI 对话与语音唤醒"],["学习模式","口语陪练、阅读、单词、课程表、学情分析"],["显示形态","单屏菜单式交互；双屏主菜单与副屏协同显示"],["本地硬件","ESP32-S3、存储、电子纸驱动、按键与音频链路"],["云端协同","用户、内容、学习记录和模型服务按需连接"]],[1.55,4.9])
add_text("S3 系列适合需要自然语言交互、学习反馈和多步骤任务的场景，也可扩展到健康管理、家庭语音入口与企业定制终端。")

# 6
page(); title_page("05  NFC 蓝牙挂件与全彩系列","从随身近场到大尺寸空间展示","不同连接方式与色彩能力，让信息出现在更合适的位置。")
tbl=doc.add_table(rows=1,cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
for c in tbl.rows[0].cells: set_cell_border(c); margins(c,120,120,120,120)
tbl.cell(0,0).paragraphs[0].add_run().add_picture(str(nfc),width=Inches(2.65))
tbl.cell(0,1).paragraphs[0].add_run().add_picture(str(e6),width=Inches(2.8))
add_text("左：1.54 英寸 NFC 场景示例    右：大尺寸全彩电子纸实机",8,False,GRAY,WD_ALIGN_PARAGRAPH.CENTER,12)
table(["系列","优势","典型用途"],[["NFC / 蓝牙挂件","近场更新、体积小、无需复杂配网","冰箱贴、钥匙扣、活动信息、随身提醒"],["7 / 10 英寸 E6 全彩","大画幅、丰富色彩、可扩展语音","艺术画框、门店展示、企业信息、家庭中控"]],[1.65,2.5,2.3])
add_text("电子纸的价值不只在省电，更在于画面能够脱离持续供电而保留。小尺寸强调便携和触达，大尺寸强调信息层级和空间融入。")

# 7
page(); title_page("06  内容模式生态","内容随生活持续生长","预设模式、AI 生成与自定义模式共同组成可持续扩展的内容供给。")
picture(daily,5.6,"每日信息模式示例")
table(["类别","代表模式"],[["每日信息","每日、天气、简报、今日历史、老黄历"],["文化与思考","诗词、斯多葛、禅意、认知偏差、每日一问"],["生活与习惯","食谱、倒计时、便签、打卡、健身、人生进度条"],["创作与互动","画廊、慢信、微故事、每日一谜、微挑战、毒舌"],["开放定制","后台自由组合文本、图片、数据源、模板与刷新规则"]],[1.6,4.85])
add_text("官网当前列出 23 个具名模式加自定义模式，共 24 类内置内容能力。不同模式可以通过统一的模板和数据接口适配多种尺寸与颜色。")

# 8
page(); title_page("07  开放平台与设备运营","从一次交付走向持续运营","统一后台把设备、内容、用户、计划和场景组织为可管理的产品体系。")
table(["平台模块","核心职责","面向对象"],[["设备管理","绑定、分组、状态、版本与远程配置","运营人员、渠道、客户管理员"],["内容后台","模板、素材、数据源与多尺寸适配","内容运营、设计与业务人员"],["模式广场","发现、预览、分享与复用模式","终端用户与创作者"],["在线刷机与配置","固件烧录、网络与设备参数配置","用户、售后与研发"],["开放接口","AI、家庭物联网和第三方业务系统连接","开发者与解决方案伙伴"]],[1.45,3.15,1.85])
add_text("端到端运行流程",12,True,NAVY,space_after=5)
bullets(["设备注册并获取屏幕尺寸、色彩能力和固件信息。","后台选择模式、数据源、刷新计划与设备分组。","内容服务生成或拉取数据，完成排版与图像处理。","设备通过安全连接获取更新，刷新电子纸并进入低功耗状态。","运营侧根据设备状态和内容效果持续调整。"])

# 9 architecture
make_architecture(); page(); title_page("08  技术架构","端云协同的五层产品架构","硬件、连接、平台、AI 内容与应用场景彼此解耦，并通过统一设备模型和内容模型协同。")
picture(ARCH,6.75)
add_text("设备侧以 ESP32-C3 / S3 为核心，管理电子纸刷新、音频、按键、存储与低功耗状态；云端负责设备管理、模式编排、内容生成和用户服务。连接层兼容 Wi-Fi、蓝牙与 NFC，并可按项目接入 Home Assistant、HTTPS、MQTT 与 OTA。",9.3,False,GRAY)

# 10
page(); title_page("09  场景解决方案","同一平台适配多种空间与人群","以常显、低打扰和可远程更新为共同基础，按场景组合硬件与内容。")
table(["场景","解决方案组成","价值"],[["家庭物联网控制台","Home Assistant + 状态卡片 + 语音入口","查看灯光、空调、窗帘与传感器，不必寻找遥控器"],["共享家庭看板","日历、待办、照片、重要提醒与远程更新","让不同代际共享关键家庭信息"],["便携学习终端","黑白快刷、单词卡、电子书、AI 图片与作业推送","降低发光屏干扰，形成专注学习界面"],["教育陪伴","S3 语音、口语陪练、阅读、课程表与学情分析","将听说练习、反馈和展示集中在一台终端"],["文化与探索","易经、紫微斗数、星座及知识卡片","以纸感方式承载日更解读和个性化内容"],["商业展示","大尺寸全彩屏、后台分组与定时刷新","降低布线与持续亮屏成本，融入门店与办公空间"]],[1.25,2.9,2.35])

# 11
page(); title_page("10  核心技术与交付能力","从开发板到量产定制","三套开发板方向与统一软件平台支持原型、试制和批量交付。")
picture(board,5.8,"双屏主板实物示例")
table(["技术方向","能力摘要"],[["硬件平台","ESP32-C3、ESP32-S3 无语音版、ESP32-S3 语音版"],["显示适配","1.54、4.2、5.83、7.3、10 英寸及多色电子纸"],["固件工程","Arduino、ESP-IDF、PlatformIO、在线刷机、OTA 与低功耗"],["内容工程","模板布局、图像抖动与色彩映射、多尺寸渲染、AI 生成"],["项目交付","产品咨询、样机测试、批量采购、软硬件定制与平台集成"]],[1.5,4.95])

# 12
page(); title_page("11  公司与团队","以电子纸产品化连接 AI 与真实空间","团队聚焦电子纸终端、嵌入式系统、AI 内容平台与场景化产品落地。")
add_text("公司主体",12,True,NAVY,space_after=4)
add_text("沄墨电子科技(四川)有限责任公司")
add_text("创始人",12,True,NAVY,space_after=4)
add_text("张少敏")
add_text("团队背景",12,True,NAVY,space_after=4)
add_text("核心团队成员背景涵盖武汉电子研究所、东南大学硕士与中国科学技术大学博士，在电子工程、嵌入式开发、人工智能与产品工程化方面形成交叉能力。")
add_text("合作方向",12,True,NAVY,space_after=4)
bullets(["小悟系列产品采购与渠道合作","教育、家庭、健康与商业展示解决方案","电子纸开发板、固件与开放平台合作","样机验证、联合研发与软硬件定制"])
picture(logo,1.15)
add_text("官方网站  https://xiaowuai.vip/zh",10,True,TEAL,WD_ALIGN_PARAGRAPH.CENTER,3)
add_text("本白皮书中的产品信息基于官网公开内容整理。具体规格、功能、价格和交付计划以正式合同及最终发布版本为准。",8.5,False,GRAY,WD_ALIGN_PARAGRAPH.CENTER,0)

# footer page numbers
for section in doc.sections:
    footer=section.footer; p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("小悟 AI 墨水屏产品白皮书  ·  "); font(8,False,GRAY)(r)
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); p._p.append(fld)

doc.core_properties.title="小悟 AI 墨水屏产品白皮书"
doc.core_properties.subject="产品目录 内容生态 技术架构与公司介绍"
doc.core_properties.author="沄墨电子科技(四川)有限责任公司"
doc.save(DOCX)
print(DOCX)
